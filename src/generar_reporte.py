from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build_report(results_dir: str = "resultados", output: str = "reporte/Reporte_Tecnico_Evaluacion_Validacion.docx") -> Path:
    r = Path(results_dir)
    summary = json.loads((r / "resumen.json").read_text(encoding="utf-8"))
    ab = json.loads((r / "simulacion_ab_sintetica.json").read_text(encoding="utf-8"))
    metrics = pd.read_csv(r / "metricas_modelos.csv")
    cv = pd.read_csv(r / "validacion_cruzada.csv")
    thresholds = pd.read_csv(r / "analisis_umbral.csv")
    validation_thresholds = pd.read_csv(r / "validacion_umbral.csv")
    classes = pd.read_csv(r / "reporte_clases_logistica.csv", index_col=0)
    d, s = summary["diagnostics"], summary["statistics"]
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.5); sec.bottom_margin = Inches(.5)
    sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(8.5)

    def title(text): doc.add_heading(text, level=1)
    def para(text): doc.add_paragraph(text)
    def hyperlink(paragraph, text, url):
        relationship_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
        properties.append(color); properties.append(underline); run.append(properties)
        text_node = OxmlElement("w:t"); text_node.text = text; run.append(text_node)
        link.append(run); paragraph._p.append(link)
    def table(frame, cols=None, digits=3):
        if cols: frame = frame[cols]
        t = doc.add_table(rows=1, cols=len(frame.columns)); t.style = "Light Shading Accent 1"
        for c, name in zip(t.rows[0].cells, frame.columns): c.text = str(name)
        for _, row in frame.iterrows():
            cells = t.add_row().cells
            for c, val in zip(cells, row):
                c.text = f"{val:.{digits}f}" if isinstance(val, (float,)) else str(val)
        return t
    def page(): doc.add_page_break()

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("MAESTRÍA EN INTELIGENCIA ARTIFICIAL\n\n").bold = True
    p.add_run("Evaluar y validar modelos mediante métricas y pruebas\n\n").bold = True
    p.add_run("Clasificación de sentimiento en comentarios sobre aerolíneas\n\n\n")
    p.add_run("Reporte técnico\n\n")
    p.add_run("Fecha de ejecución: 24 de octubre de 2026\n")
    p.add_run("Dataset: Twitter US Airline Sentiment")
    dashboard = doc.add_paragraph()
    dashboard.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dashboard.add_run("Tablero visual de resultados: ")
    hyperlink(dashboard, "Abrir tablero en línea", "https://papila9.github.io/actividad6-evaluacion-validacion/resultados/tablero_resultados.html")
    page()

    title("1. Contexto, problema y objetivo")
    para("Las opiniones en redes sociales pueden ayudar a priorizar la revisión de comentarios de clientes. Este ejercicio evalúa modelos que clasifican texto en sentimiento negativo, neutral o positivo. El objetivo no es desplegar un sistema operativo, sino comparar modelos con particiones que reduzcan la fuga causada por textos repetidos.")
    para("Objetivo SMART: comparar, sobre una partición agrupada y estratificada reproducible, un modelo base, Regresión Logística y Random Forest; cuantificar desempeño por clase, estabilidad en cinco folds y capacidad de priorización de negativos, generando métricas y artefactos reproducibles con semilla 42.")
    title("2. Datos y preparación")
    para(f"El archivo Twitter US Airline Sentiment [1] contiene {d['records_before']:,} filas y {len(d['columns'])} columnas. La variable de texto es `text` y la etiqueta es `airline_sentiment`. Se registraron tipos de datos y valores faltantes por columna en el resumen JSON; no hubo textos o etiquetas faltantes ni textos vacíos. Antes de modelar se detectaron {d['conflicting_text_groups']} textos con etiquetas contradictorias, representados en {d['conflicting_rows_excluded']} filas; esos grupos se excluyeron del modelado y se conservaron en el diagnóstico. Los duplicados con una sola etiqueta se redujeron a una fila por texto.")
    table(pd.DataFrame([{"Etapa": "Antes", "Filas": d["records_before"]}, {"Etapa": "Después de depurar", "Filas": d["records_after"]}]))
    para("La etiqueta original está desbalanceada (9,178 negativas; 3,099 neutrales; 2,363 positivas). Al retirar grupos contradictorios y duplicados exactos quedaron 14,402 textos: 9,079 negativos, 3,046 neutrales y 2,277 positivos. La eliminación de conflictos evita asignar dos objetivos incompatibles al mismo texto, pero reduce la muestra y puede cambiar su distribución.")
    page()

    title("3. Diseño experimental y modelos")
    para(f"Se usó StratifiedGroupKFold [2] con cinco folds y semilla {d['seed']}. Un fold se reservó como holdout: {d['partition']['train_rows']:,} textos para entrenamiento y {d['partition']['test_rows']:,} para prueba. Se verificaron cero textos compartidos. En la porción de entrenamiento se aplicó validación cruzada agrupada de cinco folds. TF-IDF se incluye en cada pipeline, de modo que el vocabulario se ajusta dentro de cada partición.")
    para("Se compararon DummyClassifier (clase mayoritaria), Regresión Logística con ponderación balanceada y Random Forest. Las métricas incluyen exactitud, precisión/recall macro y F1 macro/ponderado. Se prioriza F1 macro porque otorga el mismo peso a las tres clases.")
    title("4. Desempeño en el holdout")
    table(metrics.rename(columns={"model":"Modelo","accuracy":"Exactitud","precision_macro":"Precisión macro","recall_macro":"Recall macro","f1_macro":"F1 macro","f1_weighted":"F1 ponderado"}), digits=3)
    chosen_name = summary["best_by_macro_f1"]
    chosen_cv = cv.loc[cv.model == chosen_name].iloc[0]
    chosen_test = metrics.loc[metrics.model == chosen_name].iloc[0]
    para(f"El modelo se seleccionó por F1 macro medio de validación cruzada en los datos de entrenamiento. {chosen_name} obtuvo la mayor media ({chosen_cv['f1_macro_mean']:.3f}, DE {chosen_cv['f1_macro_std']:.3f}) y después alcanzó F1 macro {chosen_test['f1_macro']:.3f} y exactitud {chosen_test['accuracy']:.3f} en el holdout reservado. Random Forest obtuvo F1 macro de prueba {metrics.loc[metrics.model == 'Random Forest', 'f1_macro'].iloc[0]:.3f}; DummyClassifier, {metrics.loc[metrics.model == 'DummyClassifier', 'f1_macro'].iloc[0]:.3f}.")
    title("5. Desempeño por clase")
    table(classes.loc[["negative", "neutral", "positive"]].reset_index(names="Clase").rename(columns={"index":"Clase","precision":"Precisión","recall":"Recall","f1-score":"F1","support":"Casos"}))
    para("En la matriz de Regresión Logística, 1,569 negativos se reconocieron correctamente y 189 se confundieron con neutrales. De 610 neutrales, 135 se confundieron con negativos. Para la clase positiva, 67 de 455 se predijeron como neutrales. El F1 neutral (0.642) es el más bajo; los errores más frecuentes se producen entre negativo y neutral.")
    doc.add_picture(str(r / "matriz_confusion_logistica.png"), width=Inches(4.4))
    page()

    title("5.1. Matriz y métricas de Random Forest")
    rf_matrix = pd.read_csv(r / "matriz_confusion_random_forest.csv", index_col=0)
    rf_class = pd.read_csv(r / "reporte_clases_random_forest.csv", index_col=0)
    para("La matriz siguiente usa el mismo conjunto de prueba y el mismo orden de clases que la regresión logística.")
    table(rf_matrix.reset_index(names="Real"))
    table(rf_class.loc[["negative", "neutral", "positive"]].reset_index(names="Clase").rename(columns={"precision":"Precisión","recall":"Recall","f1-score":"F1","support":"Casos"}))
    para("Random Forest clasificó correctamente 1,478 negativos, 400 neutrales y 306 positivos. Confundió 243 negativos con neutrales y 146 neutrales con negativos. Su desempeño por clase fue menor que el de Regresión Logística en este holdout; ambas matrices están disponibles en formato PNG y CSV.")
    doc.add_picture(str(r / "matriz_confusion_random_forest.png"), width=Inches(4.0))
    page()

    title("6. Validación cruzada y ROC")
    table(cv.rename(columns={"model":"Modelo","f1_macro_mean":"F1 macro medio","f1_macro_std":"Desv. estándar","f1_macro_min":"Mínimo","f1_macro_max":"Máximo","accuracy_mean":"Exactitud media"}),
          ["Modelo","F1 macro medio","Desv. estándar","Mínimo","Máximo","Exactitud media"], digits=3)
    lr_cv = cv.loc[cv.model == "Regresión logística"].iloc[0]
    rf_cv = cv.loc[cv.model == "Random Forest"].iloc[0]
    para(f"El F1 macro medio fue {lr_cv['f1_macro_mean']:.3f} (DE {lr_cv['f1_macro_std']:.3f}; rango {lr_cv['f1_macro_min']:.3f}–{lr_cv['f1_macro_max']:.3f}) para Regresión Logística y {rf_cv['f1_macro_mean']:.3f} (DE {rf_cv['f1_macro_std']:.3f}; rango {rf_cv['f1_macro_min']:.3f}–{rf_cv['f1_macro_max']:.3f}) para Random Forest. El orden se mantuvo estable entre folds. Esto describe estabilidad bajo estas particiones, no garantiza generalización a otras fuentes o periodos.")
    para(f"El AUC macro OvR fue {s['auc_ovr_macro_logistic']:.3f} para Regresión Logística y {s['auc_ovr_macro_random_forest']:.3f} para Random Forest. El valor resume la capacidad de ordenar positivos frente a negativos para cada clase en este conjunto; no sustituye precisión, recall ni calibración.")
    doc.add_picture(str(r / "curvas_roc.png"), width=Inches(6.0))
    page()
    title("7. Análisis de umbral negativo")
    recommended = float(s["recommended_negative_threshold"])
    val_choice = validation_thresholds.loc[validation_thresholds.threshold == recommended].iloc[0]
    test_choice = thresholds.loc[thresholds.threshold == recommended].iloc[0]
    table(validation_thresholds.rename(columns={"threshold":"Umbral de validación","precision":"Precisión","recall":"Recall","f1":"F1 validación","false_positives":"FP de validación","false_negatives":"FN de validación"}), digits=3)
    table(thresholds.rename(columns={"threshold":"Umbral","precision":"Precisión","recall":"Recall","f1":"F1","false_positives":"Falsos positivos","false_negatives":"Falsos negativos"}), digits=3)
    para(f"El umbral recomendado ({recommended:.2f}) se eligió maximizando F1 binario sobre probabilidades out-of-fold del entrenamiento; su F1 de validación fue {val_choice['f1']:.3f}. Fijado ese valor antes de consultar el holdout, la evaluación final obtuvo precisión {test_choice['precision']:.3f}, recall {test_choice['recall']:.3f}, {int(test_choice['false_positives'])} alertas falsas y {int(test_choice['false_negatives'])} negativos omitidos. La tabla de umbrales del holdout es descriptiva y no se usó para escoger el recomendado. Más recall implica mayor carga de revisión; más precisión puede dejar negativos sin detectar.")
    doc.add_picture(str(r / "analisis_umbral.png"), width=Inches(4.3))
    page()

    title("8. Comparación emparejada")
    table(pd.DataFrame([
        {"Random Forest": "Correcto", "Regresión logística correcta": s["contingency_correct_both"], "Regresión logística incorrecta": s["contingency_forest_only"]},
        {"Random Forest": "Incorrecto", "Regresión logística correcta": s["contingency_logistic_only"], "Regresión logística incorrecta": s["contingency_both_wrong"]},
    ]))
    para(f"La tabla de errores emparejados mostró {s['contingency_logistic_only']} casos correctos solo para Regresión Logística y {s['contingency_forest_only']} correctos solo para Random Forest. McNemar [6] con corrección de continuidad reportó χ²={s['mcnemar_chi2_cc']:.3f}; p exacta binomial={s['mcnemar_exact_p']:.3g}. Con α=0.05 se rechaza igualdad de tasas marginales de error en este conjunto. Es una comparación de errores emparejados, no una prueba A/B operativa.")
    title("9. Diseño y simulación A/B")
    para("SIMULACIÓN CON DATOS SINTÉTICOS. Se fijaron recall de referencia 0.65, recall asistido 0.75, α=0.05, potencia objetivo 80% y semilla 42. El cálculo aproximado exige 329 casos negativos por brazo. La simulación alcanzó recalls observados de {ab['groups'][0]['recall']:.3f} y {ab['groups'][1]['recall']:.3f}; diferencia {ab['test']['difference_b_minus_a']:.3f}, z={ab['test']['z_statistic']:.3f}, p={ab['test']['p_value']:.3f}, IC95% [{ab['test']['confidence_interval_95'][0]:.3f}, {ab['test']['confidence_interval_95'][1]:.3f}].")
    para("La diferencia simulada no fue significativa a α=0.05. Es un resultado sintético con supuestos editables, no evidencia de mejora real. El tamaño calculado no garantiza validez: la prueba real requiere asignación adecuada, independencia, casos suficientes, etiquetas de referencia consistentes y protocolo previo.")
    doc.add_picture(str(r / "simulacion_ab_sintetica.png"), width=Inches(4.4))
    page()

    title("10. Impacto potencial, limitaciones y conclusiones")
    impact = summary["business_impact"]
    para(f"Proyección ilustrativa para {impact['volume']:,} comentarios con prevalencia del holdout y umbral {impact['threshold']:.2f}: aproximadamente {impact['projected_negative_detected']} negativos detectados, {impact['projected_negative_missed']} negativos no detectados y {impact['projected_false_alerts']} alertas falsas; cerca de {impact['projected_alerts_for_review']} alertas requerirían revisión. Son estimaciones escaladas desde métricas observadas, no conteos de operación. No se monetiza el impacto porque faltan costos de revisión, costo de omisiones y capacidad del equipo.")
    para("Limitaciones: mensajes de 2015 y dominio específico; resolución de conflictos mediante exclusión; evaluación sobre un único holdout; ausencia de evaluación temporal o externa; resultados sensibles a la prevalencia; A/B sintética basada en supuestos; falta de datos operativos de tiempo y costos. El análisis de una prueba A/B real queda pendiente porque no se proporcionó un CSV operacional; la simulación solo demuestra el procedimiento.")
    para(f"Conclusión: {chosen_name} fue el mejor de los modelos comparados por F1 macro medio en la validación cruzada y mantuvo un desempeño superior en el holdout reservado. Los errores neutrales siguen siendo relevantes. Antes de utilizar el modelo para priorizar atención, se recomienda evaluar datos actuales y representativos, definir el costo de omisiones frente al volumen de alertas y realizar una prueba operacional preespecificada.")
    title("Referencias")
    for ref in [
        "[1] CrowdFlower. Twitter US Airline Sentiment. Kaggle. https://www.kaggle.com/datasets/crowdflower/twitter-airline-sentiment",
        "[2] Scikit-learn Developers. Cross-validation: evaluating estimator performance. https://scikit-learn.org/stable/modules/cross_validation.html",
        "[3] Dietterich, T. G. (1998). Approximate Statistical Tests for Comparing Supervised Classification Learning Algorithms. Neural Computation, 10(7), 1895–1923. https://doi.org/10.1162/089976698300017197",
        "[4] Fawcett, T. (2006). An introduction to ROC analysis. Pattern Recognition Letters, 27(8), 861–874. https://doi.org/10.1016/j.patrec.2005.10.010",
        "[5] Cohen, J. (1988). Statistical Power Analysis for the Behavioral Sciences (2nd ed.). Lawrence Erlbaum Associates.",
        "[6] McNemar, Q. (1947). Note on the sampling error of the difference between correlated proportions or percentages. Psychometrika, 12(2), 153–157. https://doi.org/10.1007/BF02295996"
    ]: doc.add_paragraph(ref, style="List Bullet")
    output_path = Path(output); output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_report())
